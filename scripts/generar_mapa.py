#!/usr/bin/env python3
"""
Reads SORTIS_FLOOR_PLAN.xlsx and PTY Users & Roles.xlsx (Query - Active sheet),
then outputs positions.json to the OneDrive FLOOR PLAN folder for the SORTIS Desktop app.
"""

import json
import os
import sys
from datetime import datetime, timezone

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from src.core.path_config import get_floor_plan_path


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(SCRIPT_DIR, "..", "SORTIS_FLOOR_PLAN.xlsx")
PTY_USERS_PATH = r"C:\Users\david.arjona\OneDrive - a360inc\PTY Files - PTY Users\PTY Users & Roles.xlsx"


def fix_encoding(s):
    if not isinstance(s, str):
        return s
    try:
        return s.encode("latin-1").decode("utf-8")
    except (UnicodeDecodeError, UnicodeEncodeError):
        return s


def normalize_seat_no(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return str(int(val))
    return str(val).strip()


def read_active_users():
    if not os.path.exists(PTY_USERS_PATH):
        print(f"  WARNING: PTY Users file not found at {PTY_USERS_PATH}")
        return {}
    wb = openpyxl.load_workbook(PTY_USERS_PATH, read_only=True, data_only=True)
    ws = wb["Query - Active"]
    people = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        emp_no = row[0]
        title = row[4]
        dept = row[5]
        full_name = row[7]
        if not full_name:
            continue
        name = fix_encoding(str(full_name).strip())
        dept_str = fix_encoding(str(dept).strip()) if dept else None
        title_str = fix_encoding(str(title).strip()) if title else None
        emp_no_str = str(int(emp_no)) if emp_no else None
        people[name] = {
            "department": dept_str,
            "title": title_str,
            "employee_number": emp_no_str,
        }
    wb.close()
    return people


def read_seats_allocation(wb):
    ws = wb["Seats Allocation"]
    seats = {}
    for row in range(2, ws.max_row + 1):
        seat_no_raw = ws.cell(row, 1).value
        name = ws.cell(row, 2).value
        brigadista = ws.cell(row, 3).value
        notas = ws.cell(row, 4).value
        department = ws.cell(row, 5).value
        title = ws.cell(row, 6).value
        status = ws.cell(row, 7).value
        name_str = fix_encoding(str(name).strip()) if name else None
        seat_key = normalize_seat_no(seat_no_raw)
        if seat_key is None:
            continue
        seats[seat_key] = {
            "name": name_str if name_str and name_str != "-" else None,
            "brigadista": fix_encoding(str(brigadista).strip()) if brigadista else None,
            "notas": fix_encoding(str(notas).strip()) if notas else None,
            "department": fix_encoding(str(department).strip()) if department else None,
            "title": fix_encoding(str(title).strip()) if title else None,
            "status": fix_encoding(str(status).strip().upper()) if status else None,
        }
    return seats


def read_room_names(wb):
    """Read the 'Room Names' worksheet and return a dict keyed by geometry."""
    if "Room Names" not in wb.sheetnames:
        return {}
    ws = wb["Room Names"]
    overrides = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        try:
            mr = int(row[0])
            MR = int(row[1])
            mc = int(row[2])
            MC = int(row[3])
        except (TypeError, ValueError, IndexError):
            continue
        name = fix_encoding(str(row[4]).strip()) if row[4] is not None else ""
        if name:
            overrides[(mr, MR, mc, MC)] = name
    return overrides


def read_floor_plan(wb, room_overrides=None):
    ws = wb["Floor plan"]
    merged_map = {}
    rooms = []
    for mc in ws.merged_cells.ranges:
        val = ws.cell(mc.min_row, mc.min_col).value
        if val is not None:
            val = fix_encoding(str(val).strip())
        geo = (mc.min_row, mc.max_row, mc.min_col, mc.max_col)
        override = (room_overrides or {}).get(geo)
        name = override or val or ""
        info = {
            "name": name,
            "min_row": mc.min_row,
            "min_col": mc.min_col,
            "max_row": mc.max_row,
            "max_col": mc.max_col,
        }
        for r in range(mc.min_row, mc.max_row + 1):
            for c in range(mc.min_col, mc.max_col + 1):
                merged_map[(r, c)] = info
        rooms.append(info)

    seats = []
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            key = (row, col)
            if key in merged_map:
                continue
            val = ws.cell(row, col).value
            if val is not None:
                try:
                    seat_num = int(val)
                    seats.append({
                        "row": row,
                        "col": col,
                        "seat_no": str(seat_num),
                    })
                except (ValueError, TypeError):
                    pass

    max_row = 1
    max_col = 1
    for info in rooms:
        max_row = max(max_row, info["max_row"])
        max_col = max(max_col, info["max_col"])
    for s in seats:
        max_row = max(max_row, s["row"])
        max_col = max(max_col, s["col"])

    return rooms, seats, max_row, max_col


def build_output(rooms, grid_seats, seats_alloc, active_users):
    departments = set()
    output_seats = []
    for gs in grid_seats:
        seat_no = gs["seat_no"]
        alloc = seats_alloc.get(seat_no, {})
        person_name = alloc.get("name")
        person = None
        if person_name:
            pty = active_users.get(person_name, {})
            dept = pty.get("department") or alloc.get("department")
            if dept:
                departments.add(dept)
            person = {
                "name": person_name,
                "department": dept,
                "title": pty.get("title") or alloc.get("title"),
                "employee_number": pty.get("employee_number"),
            }
        status = alloc.get("status")
        output_seats.append({
            "seat_no": seat_no,
            "row": gs["row"],
            "col": gs["col"],
            "person": person,
            "brigadista": alloc.get("brigadista"),
            "notas": alloc.get("notas"),
            "status": status if status in ("ERROR",) else None,
        })

    return {
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "grid": {"rows": max(s["row"] for s in grid_seats) if grid_seats else 0,
                 "cols": max(s["col"] for s in grid_seats) if grid_seats else 0},
        "rooms": rooms,
        "seats": output_seats,
        "departments": sorted(departments),
    }


def main():
    excel_path = os.path.abspath(EXCEL_PATH)
    if not os.path.exists(excel_path):
        print(f"ERROR: Excel not found at {excel_path}")
        sys.exit(1)

    print(f"Reading: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    room_overrides = read_room_names(wb)
    print(f"  Room name overrides: {len(room_overrides)}")
    rooms, grid_seats, max_row, max_col = read_floor_plan(wb, room_overrides)
    print(f"  Rooms: {len(rooms)}, Seats on grid: {len(grid_seats)}")
    print(f"  Grid: {max_row}x{max_col}")

    seats_alloc = read_seats_allocation(wb)
    wb.close()
    print(f"  Seats allocation entries: {len(seats_alloc)}")

    print(f"Reading Query - Active...")
    active_users = read_active_users()
    print(f"  Active users entries: {len(active_users)}")

    data = build_output(rooms, grid_seats, seats_alloc, active_users)
    data["grid"]["rows"] = max_row
    data["grid"]["cols"] = max_col

    output_dir = get_floor_plan_path()
    if not output_dir:
        print("ERROR: FLOOR PLAN directory not found via OneDrive path.")
        sys.exit(1)
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "positions.json")
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"\nWritten: {output_path}")
    print(f"  Seats: {len(data['seats'])}, Departments: {len(data['departments'])}")


if __name__ == "__main__":
    main()
