"""Excel persistence for room name overrides.

The 'Room Names' worksheet in SORTIS_FLOOR_PLAN.xlsx stores admin-defined
room name overrides. Each entry is keyed by room geometry (min_row, max_row,
min_col, max_col) which matches the Floor plan sheet and positions.json.

Worksheet structure:
    A: min_row   (int)
    B: max_row   (int)
    C: min_col   (int)
    D: max_col   (int)
    E: name      (str, non-empty to apply)
"""

import os

import openpyxl

from src.core.path_config import get_floor_plan_path

_SHEET_NAME = "Room Names"
_HEADERS = ["min_row", "max_row", "min_col", "max_col", "name"]


def get_excel_path():
    """Return the full path to SORTIS_FLOOR_PLAN.xlsx, or None."""
    base = get_floor_plan_path()
    if not base:
        return None
    path = os.path.join(base, "SORTIS_FLOOR_PLAN.xlsx")
    return path if os.path.exists(path) else None


def _geometry_key(min_row, max_row, min_col, max_col):
    return (min_row, max_row, min_col, max_col)


def read_room_names(excel_path):
    """Read the Room Names worksheet and return a dict keyed by geometry tuple.

    Returns:
        dict[(min_row, max_row, min_col, max_col) -> str]: name overrides.
        Only entries with non-empty names are included.
    """
    if not excel_path or not os.path.exists(excel_path):
        return {}
    try:
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    except (openpyxl.utils.InvalidFileException, OSError):
        return {}
    if _SHEET_NAME not in wb.sheetnames:
        wb.close()
        return {}
    ws = wb[_SHEET_NAME]
    result = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        try:
            mr = int(row[0])
            MR = int(row[1])
            mc = int(row[2])
            MC = int(row[3])
        except (TypeError, ValueError, IndexError):
            continue
        name = str(row[4]).strip() if row[4] is not None else ""
        if name:
            result[_geometry_key(mr, MR, mc, MC)] = name
    wb.close()
    return result


def write_room_name(excel_path, min_row, max_row, min_col, max_col, name):
    """Write or update a single room name entry in the Room Names worksheet.

    Creates the worksheet if it does not exist. Saves the workbook after write.

    Raises:
        OSError: if the workbook cannot be opened or saved.
        PermissionError: if the file is locked by another process.
    """
    if not excel_path:
        raise OSError("Excel path not provided")
    wb = openpyxl.load_workbook(excel_path)
    if _SHEET_NAME in wb.sheetnames:
        ws = wb[_SHEET_NAME]
    else:
        ws = wb.create_sheet(_SHEET_NAME)
        ws.append(_HEADERS)

    key = _geometry_key(min_row, max_row, min_col, max_col)
    found = False
    for row_idx in range(2, ws.max_row + 1):
        try:
            r = int(ws.cell(row_idx, 1).value)
            R = int(ws.cell(row_idx, 2).value)
            c = int(ws.cell(row_idx, 3).value)
            C = int(ws.cell(row_idx, 4).value)
        except (TypeError, ValueError):
            continue
        if (r, R, c, C) == key:
            ws.cell(row_idx, 5).value = name
            found = True
            break
    if not found:
        ws.append([min_row, max_row, min_col, max_col, name])

    wb.save(excel_path)
    wb.close()


def room_names_from_positions(positions_data):
    """Extract room geometry dict from positions.json data for validation."""
    rooms = {}
    for r in positions_data.get("rooms", []):
        key = _geometry_key(r["min_row"], r["max_row"], r["min_col"], r["max_col"])
        rooms[key] = r.get("name", "")
    return rooms
